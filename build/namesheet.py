"""
Builds the workbook that asks a human to name the handles nobody can name.

The point of this file is that the question is ANSWERABLE. A bare list of
nicknames is not: "who is Scanc?" has no answer on its own. What makes it
answerable is the company they keep - the nights they played and the handles
they played alongside - plus a dropdown of everyone already known, so the
common case is two clicks and no typing.

The dropdown deliberately does NOT block free text. Some of these are people
who have never appeared in the roster at all, and a validation rule that
refuses their name would make the sheet unanswerable for exactly the cases
that matter most.

Run:  python build/namesheet.py
In:   importer/live/unknown-handles.json  (written by every emit run)
Out:  /mnt/user-data/outputs/MIP-unknown-nicknames.xlsx
"""
import json
import pathlib
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = pathlib.Path(__file__).resolve().parents[1]
UNKNOWNS = json.loads((ROOT / "importer/live/unknown-handles.json").read_text())
PLAYERS = json.loads(pathlib.Path("/tmp/players.json").read_text())
OUT = pathlib.Path("/mnt/user-data/outputs/MIP-unknown-nicknames.xlsx")

ARIAL = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="0A3B2C")
FILL_ME = PatternFill("solid", fgColor="FFF6C2")
EXAMPLE_FILL = PatternFill("solid", fgColor="EFF2EE")


def is_filler(handle: str) -> bool:
    """Scoreholio's own filler for an unregistered seat, not a person.

    Three shapes turn up: "Player 7 P.", a handle that is only digits
    ("12 1.", "5 5."), and repeated-letter placeholders ("Aaa B.").
    """
    words = re.sub(r"[.’']", " ", handle).split()
    if not words:
        return True
    if re.match(r"^player\b", handle, re.I):
        return True
    if all(re.fullmatch(r"\d+", w) for w in words):
        return True
    if len(words) > 1 and all(re.fullmatch(r"([a-z])\1*", w, re.I) for w in words):
        return True
    return False


def suggestion(h: dict) -> str:
    if h.get("candidates"):
        return "either " + " or ".join(h["candidates"])
    handle = h["handle"]
    # A handle that is already a first name and a surname is almost certainly
    # that person, just never entered in the roster. Say so; do not assume it.
    words = [w for w in re.sub(r"[^A-Za-z' -]", " ", handle).split() if len(w) > 1]
    if len(words) >= 2 and all(w[0].isupper() for w in words[:2]):
        return "reads like a real name already - confirm the spelling"
    return ""


handles = UNKNOWNS["handles"]
real = [h for h in handles if not is_filler(h["handle"])]
filler = [h for h in handles if is_filler(h["handle"])]

wb = Workbook()

# ----------------------------------------------------------------- fill me in
ws = wb.active
ws.title = "Nicknames to name"

ws["A1"] = "Merrick In A Pickle - nicknames the site cannot match to a person"
ws["A1"].font = Font(name=ARIAL, size=14, bold=True)
ws["A2"] = (
    "Fill in column H only. Pick a name from the dropdown, or just type one if the person "
    "is not in the list yet. Leave a row blank if you do not know - a blank is better than a guess. "
    "Type  IGNORE  if the nickname is not a real player. Send the file back and every game "
    "belonging to a named person goes onto the site."
)
ws["A2"].font = Font(name=ARIAL, size=10, italic=True, color="64756D")
ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
ws.merge_cells("A2:I2")
ws.row_dimensions[2].height = 42

headers = [
    "Nickname", "Games", "Nights", "First seen", "Last seen",
    "Played alongside", "Best guess", "REAL NAME - fill this in", "Notes (optional)",
]
for i, htext in enumerate(headers, start=1):
    c = ws.cell(row=4, column=i, value=htext)
    c.font = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
    c.fill = HEAD_FILL
    c.alignment = Alignment(vertical="center", wrap_text=True)
ws.row_dimensions[4].height = 28

# One example row, so the expected format is never in doubt.
example = ["(example) Chasette Gold", 22, 2, "2026-01-01", "2026-09-11",
           "Paul H., Andy, Bill R.", "", "Danielle Farruggia", "already solved - shown as a sample"]
for i, v in enumerate(example, start=1):
    c = ws.cell(row=5, column=i, value=v)
    c.font = Font(name=ARIAL, size=10, italic=True, color="64756D")
    c.fill = EXAMPLE_FILL

row = 6
for h in real:
    ws.cell(row=row, column=1, value=h["handle"])
    ws.cell(row=row, column=2, value=h["games"])
    ws.cell(row=row, column=3, value=h["nights"])
    ws.cell(row=row, column=4, value=h["firstSeen"])
    ws.cell(row=row, column=5, value=h["lastSeen"])
    ws.cell(row=row, column=6, value=", ".join(p["handle"] for p in h["playedWith"][:6]))
    ws.cell(row=row, column=7, value=suggestion(h))
    ws.cell(row=row, column=8).fill = FILL_ME
    for col in range(1, 10):
        ws.cell(row=row, column=col).font = Font(name=ARIAL, size=10)
    row += 1

last = row - 1

ws.freeze_panes = "A5"
for col, width in zip("ABCDEFGHI", [24, 8, 8, 12, 12, 46, 38, 30, 28]):
    ws.column_dimensions[col].width = width

# ------------------------------------------------------------- known players
names = wb.create_sheet("Known players")
names["A1"] = "Every player already on the site. This is the dropdown's source list."
names["A1"].font = Font(name=ARIAL, size=10, bold=True)
for i, n in enumerate(PLAYERS, start=2):
    c = names.cell(row=i, column=1, value=n)
    c.font = Font(name=ARIAL, size=10)
names.column_dimensions["A"].width = 34

src = f"'Known players'!$A$2:$A${len(PLAYERS) + 1}"
dv = DataValidation(type="list", formula1=src, allow_blank=True, showDropDown=False)
# Deliberately non-blocking: several of these people are not in the list at all.
dv.showErrorMessage = False
dv.prompt = "Pick a name, or type a new one. IGNORE = not a real player."
dv.promptTitle = "Real name"
dv.showInputMessage = True
ws.add_data_validation(dv)
dv.add(f"H6:H{last}")

# ------------------------------------------------------------- filler seats
fl = wb.create_sheet("Filler seats (excluded)")
fl["A1"] = "Scoreholio's placeholders for unregistered seats - excluded on purpose."
fl["A1"].font = Font(name=ARIAL, size=12, bold=True)
fl["A2"] = ("These are not people, so their games are not counted. Listed here only so the "
            "judgement is visible. If any of these IS a real player, say so and it moves to sheet 1.")
fl["A2"].font = Font(name=ARIAL, size=10, italic=True, color="64756D")
fl.merge_cells("A2:D2")
for i, htext in enumerate(["Nickname", "Games", "Nights", "First seen"], start=1):
    c = fl.cell(row=4, column=i, value=htext)
    c.font = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
    c.fill = HEAD_FILL
r = 5
for h in filler:
    fl.cell(row=r, column=1, value=h["handle"]).font = Font(name=ARIAL, size=10)
    fl.cell(row=r, column=2, value=h["games"]).font = Font(name=ARIAL, size=10)
    fl.cell(row=r, column=3, value=h["nights"]).font = Font(name=ARIAL, size=10)
    fl.cell(row=r, column=4, value=h["firstSeen"]).font = Font(name=ARIAL, size=10)
    r += 1
for col, width in zip("ABCD", [26, 8, 8, 12]):
    fl.column_dimensions[col].width = width

# A live tally, so progress is visible in the file itself.
ws["K4"] = "Filled in so far"
ws["K4"].font = Font(name=ARIAL, size=10, bold=True)
ws["L4"] = f'=COUNTA(H6:H{last})'
ws["L4"].font = Font(name=ARIAL, size=10)
ws["K5"] = "Still blank"
ws["K5"].font = Font(name=ARIAL, size=10, bold=True)
ws["L5"] = f'={last - 5}-COUNTA(H6:H{last})'
ws["L5"].font = Font(name=ARIAL, size=10)
ws["K6"] = "Games those blanks are holding back"
ws["K6"].font = Font(name=ARIAL, size=10, bold=True)
ws["L6"] = f'=SUMPRODUCT((H6:H{last}="")*(B6:B{last}))'
ws["L6"].font = Font(name=ARIAL, size=10)
ws.column_dimensions["K"].width = 32
ws.column_dimensions["L"].width = 10

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(f"wrote {OUT}")
print(f"  to name      {len(real)}")
print(f"  filler seats {len(filler)}")
print(f"  dropdown     {len(PLAYERS)} known players")
print(f"  games at stake {sum(h['games'] for h in real)}")
